"""Gera tests/fixtures/oil_bulletin_sample.xlsx: uma amostra do ficheiro real do histórico
do Weekly Oil Bulletin, com o mesmo formato (folhas, cabeçalhos e notas no fim), só com
algumas semanas e os países PT, ES e a média da UE (EU).

    uv run python scripts/make_sample.py Weekly_Oil_Bulletin_Prices_History_maticni_4web.xlsx

Fonte: Comissão Europeia, Weekly Oil Bulletin ("Reproduction is authorised provided the
source is acknowledged").
"""

import sys
from datetime import date, datetime
from pathlib import Path

import openpyxl

COUNTRIES = ("EU", "PT", "ES")
WEEKS = 5  # as mais recentes
TAX_SINCE = date(2022, 1, 1)  # e as mudanças de impostos desde esta data (mais o IVA todo)
TARGET = Path(__file__).parent.parent / "tests" / "fixtures" / "oil_bulletin_sample.xlsx"


def main(source: str) -> None:
    src = openpyxl.load_workbook(source, read_only=True, data_only=True)
    out = openpyxl.Workbook()
    out.remove(out.active)
    for name in ("Prices with taxes", "Prices wo taxes"):
        rows = list(src[name].iter_rows(values_only=True))
        header = rows[0]
        keep = [0] + [
            i
            for i, h in enumerate(header)
            if h
            and (
                str(h).split("_")[0] in COUNTRIES
                or (
                    h == "CTR"
                    and i + 1 < len(header)
                    and str(header[i + 1]).split("_")[0] in COUNTRIES
                )
            )
        ]
        sheet = out.create_sheet(name)
        weeks = 0
        for row in rows:
            if isinstance(row[0], datetime):
                if weeks == WEEKS:
                    continue
                weeks += 1
            sheet.append([row[i] if i < len(row) else None for i in keep])
    for name in ("VAT", "Excise duties", "Other Indirect Taxes"):
        sheet = out.create_sheet(name)
        country = None
        for row in src[name].iter_rows(values_only=True):
            if row[0]:
                country = str(row[0]).rstrip("_")
            when = row[1].date() if isinstance(row[1], datetime) else None
            header = when is None and country in (None, "CTR")
            wanted = country in COUNTRIES and (name == "VAT" or (when and when >= TAX_SINCE))
            if header or wanted:
                sheet.append(list(row[:8]))
    TARGET.parent.mkdir(parents=True, exist_ok=True)
    out.save(TARGET)
    print(TARGET)


if __name__ == "__main__":
    main(sys.argv[1])
