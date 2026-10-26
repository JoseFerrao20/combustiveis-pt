"""Leitura do ficheiro do histórico do Weekly Oil Bulletin (folhas "largas", com uma coluna
por país e produto) para linhas "longas", uma por semana, país e produto.

O formato não está documentado: estas regras vêm do próprio ficheiro e são confirmadas
pelos testes, com uma amostra do ficheiro real. Se a Comissão mudar o formato, a leitura
falha com :class:`FormatError` em vez de devolver dados errados.
"""

from __future__ import annotations

import io
import re
from collections.abc import Iterator
from dataclasses import dataclass
from datetime import date, datetime
from typing import Any

import openpyxl

PRICE_SHEETS = {"with_tax": "Prices with taxes", "without_tax": "Prices wo taxes"}
TAX_SHEETS = {"vat": "VAT", "excise": "Excise duties", "other": "Other Indirect Taxes"}
# Nas folhas dos impostos, os produtos estão sempre nesta ordem, da coluna C à H.
TAX_PRODUCTS = ("euro95", "diesel", "heating_oil", "fuel_oil_1", "fuel_oil_2", "lpg")
# Cabeçalho das folhas de preços: PT_price_with_tax_euro95, EU_price_wo_tax_LPG, …
_PRICE_COLUMN = re.compile(r"^([A-Z]{2,3})_price_(?:with|wo)_tax_([A-Za-z0-9_]+)$")
# Nas folhas dos impostos, cada país começa numa linha com o código e um "_" ("PT_").
_COUNTRY_CELL = re.compile(r"^([A-Z]{2,3})_$")


class FormatError(ValueError):
    """O ficheiro não tem o formato esperado (a Comissão pode tê-lo mudado)."""


@dataclass(frozen=True, slots=True)
class Price:
    week: date
    country: str  # código do boletim: PT, ES, … e as médias EU (UE) e EUR (zona euro)
    product: str  # euro95, diesel, heating_oil, fuel_oil_1, fuel_oil_2, lpg
    tax_status: str  # with_tax ou without_tax
    value: float  # por 1000 L (por tonelada nos fuelóleos), em euros


@dataclass(frozen=True, slots=True)
class Tax:
    tax: str  # vat (%), excise (ISP) ou other (outros impostos indiretos)
    country: str
    since: date  # em vigor a partir desta data
    product: str
    value: float  # IVA em %; os outros na moeda nacional por 1000 L (ou t)


@dataclass(frozen=True, slots=True)
class Bulletin:
    prices: list[Price]
    taxes: list[Tax]

    @property
    def latest_week(self) -> date:
        return max(p.week for p in self.prices)


def read_bulletin(data: bytes) -> Bulletin:
    """Lê o ficheiro do histórico (o conteúdo do .xlsx)."""
    try:
        workbook = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:  # zipfile.BadZipFile, KeyError, … conforme o que vier
        raise FormatError(f"Não é um ficheiro Excel válido: {exc}") from exc
    missing = [s for s in (*PRICE_SHEETS.values(), *TAX_SHEETS.values()) if s not in workbook]
    if missing:
        raise FormatError(f"Faltam as folhas {missing}; há {workbook.sheetnames}")
    prices = [
        price
        for status, sheet in PRICE_SHEETS.items()
        for price in _prices(workbook[sheet], status)
    ]
    taxes = [tax for kind, sheet in TAX_SHEETS.items() for tax in _taxes(workbook[sheet], kind)]
    if not prices:
        raise FormatError("O ficheiro não tem nenhum preço")
    return Bulletin(prices=prices, taxes=taxes)


def _as_date(value: Any) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    return None


def _number(value: Any) -> float | None:
    if isinstance(value, bool) or not isinstance(value, int | float):
        return None
    return float(value)


def _prices(sheet: Any, status: str) -> Iterator[Price]:
    rows = sheet.iter_rows(values_only=True)
    header = next(rows, ())
    columns: dict[int, tuple[str, str]] = {}
    for index, name in enumerate(header):
        match = _PRICE_COLUMN.match(str(name)) if name else None
        if match:
            columns[index] = (match.group(1), match.group(2).lower())
    if not columns:
        raise FormatError(f"A folha {sheet.title!r} não tem colunas PAÍS_price_…_produto")
    for row in rows:
        week = _as_date(row[0]) if row else None
        if week is None:
            continue  # as duas linhas de cabeçalho seguintes e as notas no fim
        for index, (country, product) in columns.items():
            value = _number(row[index]) if index < len(row) else None
            if value is not None:
                yield Price(week, country, product, status, value)


def _taxes(sheet: Any, kind: str) -> Iterator[Tax]:
    country: str | None = None
    for row in sheet.iter_rows(values_only=True):
        if not row:
            continue
        cell = row[0]
        if isinstance(cell, str) and (match := _COUNTRY_CELL.match(cell.strip())):
            country = match.group(1)
        since = _as_date(row[1]) if len(row) > 1 else None
        if country is None or since is None:
            continue
        for product, raw in zip(TAX_PRODUCTS, row[2:8], strict=False):
            # Uma célula vazia quer dizer "sem dados nesta data", não zero.
            value = _number(raw)
            if value is not None:
                yield Tax(kind, country, since, product, value)
